#!/usr/bin/env python3
"""Assert the authored slides carry real icons, not circle placeholders.

`ml._icon_or_dot` draws a real icon as an embedded picture (`<p:pic>`) when it
rasterizes, and silently substitutes a flat bright-blue ellipse of the same
footprint when it does not. Both look fine to `qa_check.py` and both extract as
no text at all, so the substitution is invisible to every other check in this
repo -- it only shows up in a render, as a slide full of identical blue dots.

Counting the two shape kinds catches it without a renderer, which matters here
because LibreOffice will not run headless on this machine.

EXPECTED_ICONS maps a slide index to the number of items its builder was given.
"""
import re
import sys

from pptx import Presentation

DECK = "deck/UG_London_2026.pptx"

EXPECTED_ICONS = {
    7: 2,    # security landscape  - ml.columns, 2 items
    9: 3,    # three use cases     - ml.columns, 3 items
    10: 4,   # use case 1 context  - ml.icon_text, 4 fields
    12: 4,   # use case 2 context
    14: 4,   # use case 3 context
    16: 5,   # conclusion          - ml.icon_text, 5 items
}


def main():
    prs = Presentation(DECK)
    slides = list(prs.slides)
    errors = []

    for idx, want in EXPECTED_ICONS.items():
        if idx >= len(slides):
            errors.append(f"slide {idx}: missing (expected {want} icons)")
            continue
        xml = slides[idx]._element.xml
        pics = xml.count("<p:pic>")
        dots = len(re.findall(r'prst="ellipse"', xml))
        if pics != want:
            errors.append(f"slide {idx}: {pics} embedded icon(s), expected {want}")
        if dots:
            errors.append(
                f"slide {idx}: {dots} blue-circle placeholder(s) -- an icon name "
                f"in tools/icons.py did not resolve, or cairosvg could not "
                f"rasterize it"
            )

    if errors:
        print(f"FAIL ({len(errors)} problem(s))")
        for e in errors:
            print("  -", e)
        return 1
    total = sum(EXPECTED_ICONS.values())
    print(f"PASS - {total} real icons across {len(EXPECTED_ICONS)} slides, no placeholders")
    return 0


if __name__ == "__main__":
    sys.exit(main())
